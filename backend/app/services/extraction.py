"""
Text extraction service — pulls plain text from various file types.
Every extractor returns a dict:
  {"text": str, "pages": int | None, "meta": dict, "error": str | None}
"""
import os
import csv
import io
import re
from pathlib import Path

from fastapi import HTTPException


MAX_TEXT_CHARS = 500_000  # ~500KB of text is plenty


# ============================================================
# Public API
# ============================================================

def extract_text(file_path: Path, extension: str, mime_type: str | None = None) -> dict:
    """
    Route to the correct extractor based on file extension.
    """
    ext = (extension or "").lower()

    try:
        if ext == ".pdf":
            return _extract_pdf(file_path)
        elif ext in (".docx",):
            return _extract_docx(file_path)
        elif ext in (".xlsx", ".xlsm"):
            return _extract_xlsx(file_path)
        elif ext in (".csv", ".tsv"):
            return _extract_csv(file_path, delimiter="\t" if ext == ".tsv" else ",")
        elif ext in (".txt", ".md", ".rst", ".log", ".py", ".js", ".ts",
                     ".jsx", ".tsx", ".html", ".css", ".json", ".yaml",
                     ".yml", ".xml", ".sql", ".sh", ".java", ".go",
                     ".rs", ".c", ".cpp", ".h", ".rb", ".php"):
            return _extract_plain_text(file_path, ext)
        elif ext == ".rtf":
            return _extract_rtf(file_path)
        elif ext in (".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"):
            return {
                "text": "",
                "pages": None,
                "meta": {"type": "image", "note": "OCR not yet implemented"},
                "error": None,
            }
        else:
            return {
                "text": "",
                "pages": None,
                "meta": {"type": "unknown"},
                "error": f"Extraction not supported for {ext}",
            }
    except HTTPException:
        raise
    except Exception as e:
        return {
            "text": "",
            "pages": None,
            "meta": {},
            "error": f"Extraction failed: {str(e)[:200]}",
        }


# ============================================================
# Extractors
# ============================================================

def _truncate(text: str) -> str:
    if len(text) > MAX_TEXT_CHARS:
        return text[:MAX_TEXT_CHARS] + "\n\n[TRUNCATED]"
    return text


def _extract_pdf(path: Path) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages_text = []
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception as e:
            text = f"[Error extracting page {i + 1}: {e}]"
        pages_text.append(text)

    full = "\n\n--- PAGE BREAK ---\n\n".join(pages_text)
    full = _truncate(full)

    return {
        "text": full,
        "pages": len(reader.pages),
        "meta": {
            "type": "pdf",
            "pages": len(reader.pages),
            "title": (reader.metadata.title if reader.metadata else None),
            "author": (reader.metadata.author if reader.metadata else None),
        },
        "error": None,
    }


def _extract_docx(path: Path) -> dict:
    from docx import Document

    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)

    # Also extract tables
    tables_text = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            tables_text.append("\n".join(rows))

    if tables_text:
        text += "\n\n--- TABLES ---\n\n" + "\n\n".join(tables_text)

    text = _truncate(text)

    return {
        "text": text,
        "pages": None,
        "meta": {
            "type": "docx",
            "paragraphs": len(paragraphs),
            "tables": len(doc.tables),
        },
        "error": None,
    }


def _extract_xlsx(path: Path) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    lines = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"=== Sheet: {sheet_name} ===")
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            row_count += 1
            if row_count > 500:
                lines.append("[TRUNCATED — more than 500 rows]")
                break
            cells = ["" if c is None else str(c) for c in row]
            if any(c.strip() for c in cells):
                lines.append(" | ".join(cells))
        lines.append("")

    text = _truncate("\n".join(lines))

    return {
        "text": text,
        "pages": len(wb.sheetnames),
        "meta": {
            "type": "xlsx",
            "sheets": wb.sheetnames,
        },
        "error": None,
    }


def _extract_csv(path: Path, delimiter: str = ",") -> dict:
    import chardet
    raw = path.read_bytes()
    enc = chardet.detect(raw[:4096]).get("encoding") or "utf-8"

    try:
        text = raw.decode(enc, errors="replace")
    except Exception:
        text = raw.decode("utf-8", errors="replace")

    # Pretty-print as table for the LLM
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = []
    for i, row in enumerate(reader):
        if i > 500:
            rows.append(["TRUNCATED — more than 500 rows"])
            break
        rows.append(row)

    if not rows:
        return {"text": "", "pages": None, "meta": {"type": "csv"}, "error": None}

    # Format as Markdown-like table
    lines = []
    header = rows[0]
    lines.append(" | ".join(header))
    lines.append("-" * 40)
    for row in rows[1:]:
        lines.append(" | ".join(row))

    text = _truncate("\n".join(lines))

    return {
        "text": text,
        "pages": None,
        "meta": {
            "type": "csv",
            "columns": len(header),
            "rows": len(rows),
        },
        "error": None,
    }


def _extract_plain_text(path: Path, ext: str) -> dict:
    import chardet
    raw = path.read_bytes()
    enc = chardet.detect(raw[:4096]).get("encoding") or "utf-8"

    try:
        text = raw.decode(enc, errors="replace")
    except Exception:
        text = raw.decode("utf-8", errors="replace")

    # HTML — strip tags
    if ext in (".html", ".htm", ".xml"):
        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(text, "html.parser")
            text = soup.get_text(separator="\n")
            # Collapse blank lines
            text = re.sub(r"\n{3,}", "\n\n", text)
        except Exception:
            pass

    text = _truncate(text)

    return {
        "text": text,
        "pages": None,
        "meta": {"type": "text", "encoding": enc, "chars": len(text)},
        "error": None,
    }


def _extract_rtf(path: Path) -> dict:
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError:
        return {
            "text": "",
            "pages": None,
            "meta": {"type": "rtf"},
            "error": "striprtf library not installed",
        }

    raw = path.read_text(encoding="utf-8", errors="replace")
    text = rtf_to_text(raw)
    text = _truncate(text)

    return {
        "text": text,
        "pages": None,
        "meta": {"type": "rtf"},
        "error": None,
    }